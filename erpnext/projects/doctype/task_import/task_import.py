# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt

from io import BytesIO

import frappe
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
import re

from frappe.model.document import Document

class TaskImport(Document):
	def validate(self):
		if not self.file:
			return

		self.set("detail", [])

		file_doc = frappe.get_doc("File", {"file_url": self.file})
		content = file_doc.get_content()

		workbook = load_workbook(
			filename=BytesIO(content),
			data_only=True
		)

		worksheet = workbook.active

		# ==========================
		# Mapping Header
		# ==========================
		headers = {}

		for i, cell in enumerate(worksheet[1], start=1):
			if cell.value:
				headers[str(cell.value).strip()] = i

		required_columns = [
			"ITEM PEKERJAAN",
			"BOBOT (%)",
			"Start",
			"Finish"
		]

		for column in required_columns:
			if column not in headers:
				frappe.throw(f"Column <b>{column}</b> not found.")

		# ==========================
		# Read Excel
		# ==========================
		for row in worksheet.iter_rows(min_row=2):

			subject = row[headers["ITEM PEKERJAAN"] - 1].value
			task_weight = row[headers["BOBOT (%)"] - 1].value
			start_date = row[headers["Start"] - 1].value
			finish_date = row[headers["Finish"] - 1].value

			# Hanya import task yang lengkap
			if not (
				subject
				and task_weight is not None
				and start_date
				and finish_date
			):
				continue

			if isinstance(start_date, (int, float)):
				start_date = from_excel(start_date)

			if isinstance(finish_date, (int, float)):
				finish_date = from_excel(finish_date)
			
			if start_date > finish_date:
				frappe.throw(f"Start date <b>{start_date}</b> is greater than finish date <b>{finish_date}</b> for <b>{subject}</b>.")

			self.append("detail", {
				"subject": str(subject).strip(),
				"task_weight": task_weight,
				"exp_start_date": start_date,
				"exp_end_date": finish_date
			})

	def on_submit(self):
		frappe.enqueue(
			"erpnext.api.task.create_tasks",
			queue="short",
			docname=self.name,
			enqueue_after_commit=True
		)

	def on_cancel(self):
		frappe.enqueue(
			"erpnext.api.task.cancel_tasks",
			queue="short",
			docname=self.name,
			enqueue_after_commit=True
		)

	def create_tasks(self):
		task_map = {}

		# =====================================================
		# 1. Create semua Task terlebih dahulu
		# =====================================================

		total = len(self.detail)
		completed = 0
		for d in self.detail:

			task = frappe.new_doc("Task")
			task.project = self.project
			task.subject = d.subject
			task.exp_start_date = d.exp_start_date
			task.exp_end_date = d.exp_end_date
			task.task_weight = d.task_weight
			task.status = "Open"

			task.insert(ignore_permissions=True)
			frappe.db.commit()

			d.db_set(
				"task",
				task.name,
				update_modified=False
			)

			completed += 1

			progress = round((completed / total) * 100, 2)

			frappe.publish_realtime(
				"task_sync_progress",
				{
					"docname": self.name,
					"progress": progress,
					"completed": completed,
					"total": total
				}
			)
		# doc reload
		self.reload()

	def cancel_tasks(self):
		for d in self.detail:
			if not d.task:
				continue
			if not frappe.db.exists("Task", d.task):
				continue
			task = frappe.get_doc("Task", d.task)
			task.status = "Cancelled"
			task.save(ignore_permissions=True)